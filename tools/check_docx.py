#!/usr/bin/env python3
"""Structural check of a generated DOCX: content census, fonts, page setup.

The point is to catch the failure modes that a missing renderer would hide:
lost sections, unstyled code, missing tables, fonts that would show as tofu for
Chinese text, and code lines too wide for the text column.
"""
import io
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:  # pragma: no cover - older interpreters
    pass


def is_shaded(paragraph, fill):
    ppr = paragraph._p.find(qn("w:pPr"))
    if ppr is None:
        return False
    shd = ppr.find(qn("w:shd"))
    return shd is not None and shd.get(qn("w:fill")) == fill


def census(path, markdown):
    doc = Document(path)
    styles = {}
    code_lines = []
    longest = ("", 0)
    for paragraph in doc.paragraphs:
        styles[paragraph.style.name] = styles.get(paragraph.style.name, 0) + 1
        if is_shaded(paragraph, "F4F4F4"):
            code_lines.append(paragraph.text)
        if len(paragraph.text) > longest[1]:
            longest = (paragraph.text, len(paragraph.text))

    md = io.open(markdown, encoding="utf-8").read().splitlines()
    # 代码块里的 "# ..." 是注释，不是标题，所以统计要先跳过围栏。
    md_h1 = md_h2 = md_h3 = 0
    md_code = 0
    inside = False
    for line in md:
        if line.rstrip().startswith("```"):
            inside = not inside
            continue
        if inside:
            md_code += 1
            continue
        if line.startswith("### "):
            md_h3 += 1
        elif line.startswith("## "):
            md_h2 += 1
        elif line.startswith("# "):
            md_h1 += 1
    md_tables = 0
    for i, line in enumerate(md[:-1]):
        nxt = md[i + 1].strip()
        if line.strip().startswith("|") and "-" in nxt and \
                set(nxt) <= set("|-: "):
            md_tables += 1

    section = doc.sections[0]
    print("段落样式统计 :", ", ".join("%s=%d" % kv for kv in sorted(styles.items())))
    print("标题数       : H1=%d H2=%d H3=%d   (markdown: %d/%d/%d)"
          % (styles.get("Heading 1", 0), styles.get("Heading 2", 0),
             styles.get("Heading 3", 0), md_h1, md_h2, md_h3))
    print("表格数       : %d   (markdown: %d)" % (len(doc.tables), md_tables))
    print("代码行数     : %d   (markdown 代码行: %d)" % (len(code_lines), md_code))
    print("最长正文行   : %d 字符 -> %s" % (longest[1], longest[0][:70]))
    print("页面         : %.1f x %.1f cm, 边距 左%.1f 右%.1f 上%.1f 下%.1f"
          % (section.page_width.cm, section.page_height.cm,
             section.left_margin.cm, section.right_margin.cm,
             section.top_margin.cm, section.bottom_margin.cm))
    print("页眉/页脚    : %r / %r" % (doc.sections[0].header.paragraphs[0].text,
                                      doc.sections[0].footer.paragraphs[0].text))

    too_wide = sorted((len(line) for line in code_lines if len(line) > 96),
                      reverse=True)
    print("超过 96 字符的代码行: %d 行 %s" % (len(too_wide), too_wide[:5]))

    # Fonts: every run must carry an East Asian font, or Word renders tofu.
    missing = 0
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            rpr = run._element.find(qn("w:rPr"))
            fonts = rpr.find(qn("w:rFonts")) if rpr is not None else None
            if fonts is None or fonts.get(qn("w:eastAsia")) is None:
                missing += 1
    print("缺少 eastAsia 字体的 run: %d" % missing)

    with zipfile.ZipFile(path) as archive:
        footer = archive.read("word/footer1.xml").decode("utf-8")
    print("页脚域       : PAGE=%d NUMPAGES=%d"
          % (footer.count(" PAGE "), footer.count("NUMPAGES")))
    return 0


if __name__ == "__main__":
    sys.exit(census(sys.argv[1], sys.argv[2]))
