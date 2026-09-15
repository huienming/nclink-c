#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
# Copyright (c) 2026 huienming

"""Compare the markdown block model against what ended up in the DOCX.

Guards the renderer against silent content loss: headings, code lines and table
cells must survive the conversion unchanged.
"""
import io
import sys

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, "tools")
import md_to_docx as m  # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass


def is_code(paragraph):
    ppr = paragraph._p.find(qn("w:pPr"))
    shd = ppr.find(qn("w:shd")) if ppr is not None else None
    return shd is not None and shd.get(qn("w:fill")) == "F4F4F4"


def main(markdown, docx):
    blocks = m.parse_blocks(io.open(markdown, encoding="utf-8").read())
    doc = Document(docx)

    # 标题里也可能带行内标记（如 `code`），比较前统一去掉。
    md_headings = [m.plain(b[1]) for b in blocks
                   if b[0] in ("h1", "h2", "h3", "h4")]
    doc_headings = [p.text for p in doc.paragraphs
                    if p.style.name.startswith("Heading")]
    md_code = [line for k, b in blocks if k == "code" for line in b]
    doc_code = [p.text for p in doc.paragraphs if is_code(p)]
    md_cells = [c for k, b in blocks if k == "table"
                for row in [b[0]] + b[1] for c in row]
    doc_cells = [c.text for t in doc.tables for row in t.rows
                 for c in row.cells]

    for label, expected, actual in (
        ("标题", md_headings, doc_headings),
        ("代码", code_of(md_code), doc_code),
        ("表格单元格", [m.plain(c) for c in md_cells],
         [m.plain(c) for c in doc_cells]),
    ):
        ok = expected == actual
        print("%-10s %s (%d vs %d)" % (label, "OK" if ok else "差异",
                                       len(expected), len(actual)))
        if not ok:
            for index, (a, b) in enumerate(zip(expected, actual)):
                if a != b:
                    print("  首个差异 @%d\n    md   : %r\n    docx : %r"
                          % (index, a[:90], b[:90]))
                    break
            if len(expected) != len(actual):
                print("  长度不同，样例尾部 md=%r docx=%r"
                      % (expected[-1][:60], actual[-1][:60]))


def code_of(lines):
    """The writer turns an empty code line into a single space."""
    return [line if line else " " for line in lines]


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
