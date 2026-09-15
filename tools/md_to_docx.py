#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
# Copyright (c) 2026 huienming

"""Render MANUAL.md (or any of the project's markdown docs) as a Word document.

The markdown is parsed into a small block model which is then emitted twice:

    * to DOCX with python-docx, for handing to Word
    * to HTML, so the same content can be screenshotted for a visual check

Usage:
    python tools/md_to_docx.py MANUAL.md MANUAL.docx [--html preview.html]
"""
import argparse
import io
import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Windows ships all three; 微软雅黑 covers the Chinese text, Consolas the code.
BODY_FONT = "Calibri"
CJK_FONT = "微软雅黑"
CODE_FONT = "Consolas"


# --------------------------------------------------------------- markdown -- #

def parse_inline(text):
    """
    Split a line into [(kind, value, bold)] runs, where kind is "text" or
    "code". A **bold** span is parsed recursively, so nesting like
    ``**设备端上报的 `paths` 不带后缀**`` keeps the inner code span.
    """
    runs = []
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            runs.append(("text", text[pos:match.start()], False))
        token = match.group(0)
        if token.startswith("`"):
            runs.append(("code", token[1:-1], False))
        elif token.startswith("**"):
            for kind, value, _bold in parse_inline(token[2:-2]):
                runs.append((kind, value, True))
        else:
            label = token[1:token.index("]")]
            target = token[token.index("(") + 1:-1]
            # Internal anchors and sibling files inside the repository do not
            # gain anything from the URL being repeated in a Word document.
            if target.startswith("#") or "://" not in target:
                runs.append(("text", label, False))
            else:
                runs.append(("text", "%s (%s)" % (label, target), False))
        pos = match.end()
    if pos < len(text):
        runs.append(("text", text[pos:], False))
    return runs


def plain(text):
    """Inline text without markup, for tables of contents and previews."""
    return "".join(value for _kind, value, _bold in parse_inline(text))


def absorb_continuation(lines, i):
    """Continuation lines of a list item: indented, and not a new block."""
    parts = []
    while i < len(lines):
        nxt = lines[i]
        if not nxt.strip() or nxt[0] not in " \t":
            break
        stripped = nxt.strip()
        if stripped.startswith(("```", "|", "#", ">")):
            break
        parts.append(stripped)
        i += 1
    return parts, i


def parse_blocks(markdown):
    """Turn the markdown into a list of (kind, payload) blocks."""
    lines = markdown.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            blocks.append(("code", code))
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and \
                re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(("table", (header, rows)))
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append(("h%d" % min(level, 4), stripped[level:].strip()))
            i += 1
            continue
        if stripped == "---":
            blocks.append(("hr", None))
            i += 1
            continue
        if stripped.startswith("> "):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(quote)))
            continue
        if re.match(r"^[-*] ", stripped):
            items = []
            while i < len(lines) and re.match(r"^[-*] ", lines[i].strip()):
                parts = [lines[i].strip()[2:].strip()]
                i += 1
                more, i = absorb_continuation(lines, i)
                items.append(" ".join(parts + more))
            blocks.append(("ul", items))
            continue
        if re.match(r"^\d+\. ", stripped):
            items = []
            while i < len(lines) and re.match(r"^\d+\. ", lines[i].strip()):
                parts = [re.sub(r"^\d+\. ", "", lines[i].strip())]
                i += 1
                more, i = absorb_continuation(lines, i)
                items.append(" ".join(parts + more))
            blocks.append(("ol", items))
            continue
        if stripped == "":
            i += 1
            continue
        # paragraph: join until a blank line or a block starter
        para = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if nxt == "" or nxt.startswith(("#", "|", ">", "```", "- ", "* ")) or \
                    nxt == "---" or re.match(r"^\d+\. ", nxt):
                break
            para.append(nxt)
            i += 1
        blocks.append(("p", " ".join(para)))
    return blocks


# ------------------------------------------------------------------- docx -- #

def set_run_font(run, name, size, bold=None, color=None, cjk=None):
    run.font.name = name
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), cjk or CJK_FONT)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(paragraph, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def left_border(paragraph, color, size=18):
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), color)
    borders.append(left)
    paragraph._p.get_or_add_pPr().append(borders)


def keep_with_next(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    element = OxmlElement("w:keepNext")
    ppr.append(element)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._element.append(node)
    set_run_font(run, BODY_FONT, 9, color=RGBColor(0x60, 0x60, 0x60))


def mark_header_row(row):
    """Repeat the first row when a table spans pages (accessibility + print)."""
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_runs(paragraph, inlines, size=10.5, base_font=BODY_FONT):
    for kind, value, bold in (parse_inline(inlines) if isinstance(inlines, str)
                              else inlines):
        run = paragraph.add_run(value)
        if kind == "code":
            set_run_font(run, CODE_FONT, size - 0.5,
                         bold=bold if bold else None,
                         color=RGBColor(0xA3, 0x15, 0x15))
        else:
            set_run_font(run, base_font, size, bold=True if bold else None)


def build_docx(blocks, source_name, title):
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    headings = {
        "h1": (20, 16, 8),
        "h2": (15, 14, 6),
        "h3": (12, 10, 4),
        "h4": (10.5, 8, 3),
    }
    for kind, (size, before, after) in headings.items():
        style = doc.styles["Heading %d" % int(kind[1])]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x1F, 0x24, 0x2B)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Header and footer.
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(title)
    set_run_font(run, BODY_FONT, 8.5, color=RGBColor(0x80, 0x80, 0x80))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第 ")
    set_run_font(run, BODY_FONT, 9, color=RGBColor(0x60, 0x60, 0x60))
    add_field(footer, " PAGE ")
    run = footer.add_run(" 页 / 共 ")
    set_run_font(run, BODY_FONT, 9, color=RGBColor(0x60, 0x60, 0x60))
    add_field(footer, " NUMPAGES ")
    run = footer.add_run(" 页")
    set_run_font(run, BODY_FONT, 9, color=RGBColor(0x60, 0x60, 0x60))

    first_section = True
    for kind, payload in blocks:
        if kind == "h1":
            paragraph = doc.add_paragraph(style="Heading 1")
            if not first_section:
                paragraph.paragraph_format.page_break_before = True
            first_section = False
            add_runs(paragraph, payload, size=20)
        elif kind in ("h2", "h3", "h4"):
            style = "Heading %d" % int(kind[1])
            paragraph = doc.add_paragraph(style=style)
            size = headings[kind][0]
            add_runs(paragraph, payload, size=size)
        elif kind == "p":
            paragraph = doc.add_paragraph()
            add_runs(paragraph, payload)
        elif kind == "quote":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.4)
            paragraph.paragraph_format.space_before = Pt(2)
            left_border(paragraph, "BFBFBF")
            for rk, rv, _bold in parse_inline(payload):
                run = paragraph.add_run(rv)
                set_run_font(run, CODE_FONT if rk == "code" else BODY_FONT,
                             9.5, color=RGBColor(0x50, 0x50, 0x50))
        elif kind in ("ul", "ol"):
            # Explicit bullets and numbers: Word's list styles share one
            # numbering definition across the document, which would run 1..N
            # through unrelated lists.
            for index, item in enumerate(payload, start=1):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.85)
                paragraph.paragraph_format.first_line_indent = Cm(-0.45)
                paragraph.paragraph_format.space_after = Pt(3)
                marker = "•  " if kind == "ul" else "%d.  " % index
                run = paragraph.add_run(marker)
                set_run_font(run, BODY_FONT, 10.5)
                add_runs(paragraph, item)
        elif kind == "code":
            for line in payload:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.3)
                paragraph.paragraph_format.right_indent = Cm(0.3)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                shade(paragraph, "F4F4F4")
                run = paragraph.add_run(line if line else " ")
                set_run_font(run, CODE_FONT, 9)
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(2)
            spacer.paragraph_format.line_spacing = 1.0
        elif kind == "table":
            header, rows = payload
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table.autofit = True
            mark_header_row(table.rows[0])
            for index, cell_text in enumerate(header):
                cell = table.rows[0].cells[index]
                cell.text = ""
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(plain(cell_text))
                set_run_font(run, BODY_FONT, 9.5, bold=True)
                shade(paragraph, "EDEDED")
            for row in rows:
                cells = table.add_row().cells
                for index, cell_text in enumerate(row[:len(header)]):
                    cell = cells[index]
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    add_runs(paragraph, cell_text, size=9.5)
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(2)
        elif kind == "hr":
            continue

    return doc


# ------------------------------------------------------------------- html -- #

def escape(text):
    return (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def inline_html(text):
    out = []
    for kind, value, bold in parse_inline(text):
        if kind == "code":
            piece = "<code>%s</code>" % escape(value)
        else:
            piece = escape(value)
        out.append("<strong>%s</strong>" % piece if bold else piece)
    return "".join(out)


def build_html(blocks, title):
    out = ["<!doctype html><html><head><meta charset='utf-8'>",
           "<title>%s</title>" % escape(title),
           "<style>",
           "body{font-family:'Microsoft YaHei',Calibri,sans-serif;font-size:10.5pt;",
           "line-height:1.5;margin:24px 32px;color:#1f242b;}",
           "h1{font-size:20pt;border-bottom:2px solid #d0d0d0;padding-bottom:6px;",
           "page-break-before:always;}h1:first-of-type{page-break-before:avoid;}",
           "h2{font-size:15pt;margin-top:18px;}h3{font-size:12pt;}h4{font-size:10.5pt;}",
           "code{font-family:Consolas,monospace;color:#a31515;}",
           "pre{font-family:Consolas,monospace;font-size:9pt;background:#f4f4f4;",
           "padding:8px 10px;line-height:1.35;overflow:visible;white-space:pre-wrap;}",
           "table{border-collapse:collapse;font-size:9.5pt;margin:8px 0;}",
           "th,td{border:1px solid #999;padding:4px 6px;text-align:left;}",
           "th{background:#ededed;}", "blockquote{border-left:3px solid #bfbfbf;",
           "margin:6px 0;padding:2px 10px;color:#505050;}", "</style></head><body>"]
    for kind, payload in blocks:
        if kind in ("h1", "h2", "h3", "h4"):
            out.append("<%s>%s</%s>" % (kind, inline_html(payload), kind))
        elif kind == "p":
            out.append("<p>%s</p>" % inline_html(payload))
        elif kind == "quote":
            out.append("<blockquote>%s</blockquote>" % inline_html(payload))
        elif kind == "ul":
            out.append("<ul>" + "".join("<li>%s</li>" % inline_html(i) for i in payload) + "</ul>")
        elif kind == "ol":
            out.append("<ol>" + "".join("<li>%s</li>" % inline_html(i) for i in payload) + "</ol>")
        elif kind == "code":
            out.append("<pre>%s</pre>" % escape("\n".join(payload)))
        elif kind == "table":
            header, rows = payload
            out.append("<table><thead><tr>" +
                       "".join("<th>%s</th>" % inline_html(h) for h in header) +
                       "</tr></thead><tbody>" +
                       "".join("<tr>" + "".join("<td>%s</td>" % inline_html(c)
                                                for c in row) + "</tr>" for row in rows) +
                       "</tbody></table>")
    out.append("</body></html>")
    return "\n".join(out)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("markdown")
    parser.add_argument("docx")
    parser.add_argument("--html")
    args = parser.parse_args()

    text = io.open(args.markdown, encoding="utf-8").read()
    blocks = parse_blocks(text)
    title = blocks[0][1] if blocks and blocks[0][0] == "h1" else os.path.basename(args.markdown)
    document = build_docx(blocks, os.path.basename(args.markdown), plain(title))
    document.save(args.docx)
    print("wrote %s (%d blocks)" % (args.docx, len(blocks)))
    if args.html:
        io.open(args.html, "w", encoding="utf-8", newline="\n").write(build_html(blocks, plain(title)))
        print("wrote %s" % args.html)


if __name__ == "__main__":
    main()
